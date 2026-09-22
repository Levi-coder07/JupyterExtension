"""Append the completed low-reasoning GPT model comparison to the report."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


REPORT_PATH = Path("output/docx/linkmaker_prompt_evaluation_report.docx")
SECTION_TITLE = "Appendix D. Low-reasoning GPT model comparison"


def _set_repeat_header(row) -> None:
    """Repeat the comparison header if the table crosses a page boundary."""
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _add_row(table, values: tuple[str, ...]) -> None:
    """Add a consistently aligned comparison row."""
    cells = table.add_row().cells
    for index, (cell, value) in enumerate(zip(cells, values)):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        )


def main() -> None:
    """Append or replace the model comparison while preserving the existing report."""
    document = Document(REPORT_PATH)
    existing_heading = next(
        (paragraph for paragraph in document.paragraphs if paragraph.text.strip() == SECTION_TITLE),
        None,
    )
    if existing_heading is not None:
        body = document._element.body
        remove = False
        for child in list(body):
            if child is existing_heading._element:
                remove = True
            if remove and child.tag != qn("w:sectPr"):
                body.remove(child)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(SECTION_TITLE, level=1)
    document.add_paragraph(
        "A configuration-controlled comparison evaluated GPT-5 nano, GPT-5 mini, "
        "and GPT-5.4 mini with low reasoning against the same Titanic code-only "
        "ground truth of 36 unique Markdown-to-code pairs. Output, image, and sketch "
        "relationships were excluded. The prompt, strict schema, notebook, matching "
        "rule, and post-processing were held constant."
    )

    table = document.add_table(rows=1, cols=6)
    table.style = "Light Shading Accent 1"
    headers = (
        "Model",
        "Runs",
        "Mean links",
        "Pooled precision",
        "Pooled recall",
        "Pooled F1",
    )
    for cell, value in zip(table.rows[0].cells, headers):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_repeat_header(table.rows[0])

    _add_row(
        table,
        ("GPT-5 nano", "1", "6.0", "83.3%", "13.9%", "23.8%"),
    )
    _add_row(
        table,
        ("GPT-5 mini", "3", "37.0", "71.2%", "73.1%", "72.1%"),
    )
    _add_row(
        table,
        (
            "GPT-5.4 mini",
            "3",
            "60.3",
            "53.0%",
            "88.9%",
            "66.4%",
        ),
    )

    document.add_paragraph("Latency results")
    latency_table = document.add_table(rows=1, cols=3)
    latency_table.style = "Light Shading Accent 1"
    for cell, value in zip(
        latency_table.rows[0].cells,
        ("Model", "Mean latency", "Observed range"),
    ):
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_repeat_header(latency_table.rows[0])
    _add_row(latency_table, ("GPT-5 nano", "51.0 s", "51.0 s"))
    _add_row(latency_table, ("GPT-5 mini", "84.7 s", "76.5-92.2 s"))
    _add_row(latency_table, ("GPT-5.4 mini", "57.1 s", "32.2-84.8 s"))

    conclusion = document.add_paragraph()
    conclusion.add_run("Conclusion. ").bold = True
    conclusion.add_run(
        "GPT-5 mini produced the best balance of precision and coverage and the highest "
        "pooled F1, so it was selected as the preferred model for Markdown-to-code "
        "relationship extraction. GPT-5.4 mini achieved the highest recall but generated "
        "substantially more unmatched links, reducing precision. GPT-5 nano was precise "
        "when it returned a link but missed most ground-truth relationships at low reasoning."
    )

    limitations = document.add_paragraph()
    limitations.add_run("Interpretation limits. ").bold = True
    limitations.add_run(
        "GPT-5 nano has only one retained run in this low-reasoning model matrix, while "
        "the other models have three. The separate GPT-5 nano medium-reasoning result in "
        "Appendix C reached F1 74.7%, but it is a single configuration result and is not "
        "pooled into this table. Token usage was not retained, so this comparison reports "
        "latency but does not estimate API cost."
    )

    document.save(REPORT_PATH)


if __name__ == "__main__":
    main()
