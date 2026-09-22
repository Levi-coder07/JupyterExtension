"""Append the Markdown-to-code medium-reasoning result to the evaluation report."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK


REPORT_PATH = Path("output/docx/linkmaker_prompt_evaluation_report.docx")


def _add_metric_row(table, metric: str, value: str) -> None:
    """Add one metric/value row to an existing report table."""
    cells = table.add_row().cells
    cells[0].text = metric
    cells[1].text = value


def main() -> None:
    """Add a configuration-only evaluation appendix without altering prior results."""
    document = Document(REPORT_PATH)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading("Appendix C. Medium-reasoning configuration result", level=1)

    document.add_paragraph(
        "A configuration-only run was performed after the v2 evaluation. "
        "The Markdown-to-code prompt and strict response schema were unchanged; "
        "only the GPT-5 nano reasoning effort was changed from low to medium."
    )
    document.add_paragraph(
        "This result is evaluated directly against the same Titanic code-only "
        "ground truth: 36 unique Markdown-to-code pairs. Output, image, and "
        "sketch links are excluded from both the prediction set and the ground truth."
    )

    table = document.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Medium-reasoning result"
    _add_metric_row(table, "Predicted Markdown-to-code links", "39")
    _add_metric_row(table, "Ground-truth links", "36")
    _add_metric_row(table, "True positives", "28")
    _add_metric_row(table, "False positives", "11")
    _add_metric_row(table, "False negatives", "8")
    _add_metric_row(table, "Precision", "71.8%")
    _add_metric_row(table, "Recall / ground-truth coverage", "77.8% (28 of 36)")
    _add_metric_row(table, "F1", "74.7%")

    conclusion = document.add_paragraph()
    conclusion.add_run("Interpretation. ").bold = True
    conclusion.add_run(
        "This medium-reasoning run covers 28 of the 36 manually labeled links. "
        "It is a strong single-run result, but it should be repeated across "
        "multiple runs before replacing the reliability conclusions from the "
        "five-run v2 experiment."
    )

    reproducibility = document.add_paragraph()
    reproducibility.add_run("Configuration record. ").bold = True
    reproducibility.add_run(
        "Model: gpt-5-nano. Target: normal Markdown-to-code analysis only. "
        "Reasoning effort: medium. The independent Markdown-to-output analysis "
        "was not part of this measurement."
    )

    document.save(REPORT_PATH)


if __name__ == "__main__":
    main()
