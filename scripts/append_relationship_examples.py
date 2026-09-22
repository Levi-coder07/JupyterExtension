"""Append a readable qualitative relationship analysis to the evaluation report."""

import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPORT_PATH = Path("output/docx/linkmaker_prompt_evaluation_report.docx")
RESULT_PATH = Path("titanic.linkmaker.json")
SECTION_TITLE = "Appendix E Qualitative relationship analysis"
SELECTED_TYPES = (
    "feature_dropping",
    "age_banding_and_correlation",
    "fare_missing_value_completion",
    "display_after_fare_completion",
)


def _relationships() -> dict[str, dict[str, Any]]:
    payload = json.loads(RESULT_PATH.read_text(encoding="utf-8"))
    return {
        relationship["relationshipType"]: relationship
        for analysis in payload.get("markdownAnalyses", [])
        for relationship in analysis.get("relationships", [])
        if relationship.get("relationshipType") in SELECTED_TYPES
    }


def _remove_existing_section(document: Document) -> None:
    heading = next(
        (paragraph for paragraph in document.paragraphs if paragraph.text.strip() == SECTION_TITLE),
        None,
    )
    if heading is None:
        return
    body = document._element.body
    remove = False
    for child in list(body):
        if child is heading._element:
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def _set_cell_fill(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def _set_cell_borders(cell, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)


def _set_cell_padding(cell, value: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def _write_cell(cell, text: str, *, code: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    run.font.name = "Consolas" if code else "Arial"
    run.font.size = Pt(9 if code else 10.5)
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:ascii"), "Consolas" if code else "Arial"
    )
    run._element.get_or_add_rPr().rFonts.set(
        qn("w:hAnsi"), "Consolas" if code else "Arial"
    )
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    _set_cell_borders(cell)
    _set_cell_padding(cell)


def _add_evidence_table(document: Document, relationship: dict[str, Any]) -> None:
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Inches(2.55), Inches(3.95))
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = widths[index]

    for cell, label in zip(table.rows[0].cells, ("Markdown evidence", "Code evidence")):
        cell.text = label
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(10)
        _set_cell_fill(cell, "1F4E78")
        _set_cell_borders(cell)
        _set_cell_padding(cell)

    _write_cell(table.rows[1].cells[0], relationship["markdownPortion"]["text"])
    _write_cell(table.rows[1].cells[1], relationship["codeTarget"]["portion"]["text"], code=True)


def _add_case(
    document: Document,
    number: int,
    title: str,
    relationship: dict[str, Any],
    assessment: str,
) -> None:
    document.add_heading(f"Case {number} {title}", level=2)
    _add_evidence_table(document, relationship)

    confidence = document.add_paragraph()
    confidence.paragraph_format.space_before = Pt(7)
    confidence.add_run("Model confidence  ").bold = True
    confidence.add_run(f"{relationship['confidence'] * 100:.0f}%")

    justification = document.add_paragraph()
    justification.add_run("Model justification  ").bold = True
    justification.add_run(relationship["reason"])

    review = document.add_paragraph()
    review.add_run("Evaluation analysis  ").bold = True
    review.add_run(assessment)


def main() -> None:
    relationships = _relationships()
    missing = [item for item in SELECTED_TYPES if item not in relationships]
    if missing:
        raise ValueError(f"Missing selected relationship types: {', '.join(missing)}")

    document = Document(REPORT_PATH)
    _remove_existing_section(document)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    document.add_heading(SECTION_TITLE, level=1)
    document.add_paragraph(
        "The following cases expose the evidence presented to the evaluator and the "
        "model's own explanation. Confidence is a self-reported model estimate, not an "
        "independent correctness probability. The cases include strong, borderline, and "
        "weak links so that confidence can be interpreted together with evidence quality."
    )

    cases = (
        (
            "Directly supported transformation",
            relationships["feature_dropping"],
            "The Markdown and code name the same two columns and the same removal operation. "
            "The 98% confidence is consistent with the direct textual and operational match.",
        ),
        (
            "Supported analytical operation",
            relationships["age_banding_and_correlation"],
            "The code creates AgeBand and aggregates survival by that band, directly matching "
            "both actions in the Markdown. The selected evidence is specific and grounded.",
        ),
        (
            "Partially inconsistent implementation",
            relationships["fare_missing_value_completion"],
            "The relationship is relevant because both portions address completing Fare, but "
            "the Markdown specifies the mode while the code uses the median. A 90% confidence "
            "overstates the semantic agreement and illustrates that confidence requires review.",
        ),
        (
            "Weak contextual association",
            relationships["display_after_fare_completion"],
            "Displaying test_df does not implement the requested rounding operation. The 60% "
            "confidence appropriately signals uncertainty, but this link is better treated as "
            "a likely false positive under an evidence-based evaluation.",
        ),
    )
    for index, (title, relationship, assessment) in enumerate(cases, start=1):
        _add_case(document, index, title, relationship, assessment)

    conclusion = document.add_paragraph()
    conclusion.paragraph_format.space_before = Pt(10)
    conclusion.add_run("Overall interpretation  ").bold = True
    conclusion.add_run(
        "High confidence generally accompanies direct operation-level matches, but it does "
        "not guarantee full semantic consistency. Reviewing the Markdown, code, and reason "
        "together reveals errors that pair-level metrics and confidence alone cannot detect."
    )
    document.save(REPORT_PATH)


if __name__ == "__main__":
    main()
