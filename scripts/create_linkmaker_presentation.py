from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path("output/LinkMaker_Luna_Evaluation_Presentation.pptx")

NAVY = RGBColor(15, 28, 48)
INK = RGBColor(24, 35, 48)
MUTED = RGBColor(91, 107, 123)
TEAL = RGBColor(22, 151, 147)
ORANGE = RGBColor(235, 133, 67)
PALE = RGBColor(238, 245, 244)
PALE_BLUE = RGBColor(235, 241, 248)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(210, 220, 225)
RED = RGBColor(187, 74, 74)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.06)
    frame.margin_right = Inches(0.06)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    run = paragraph.runs[0]
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, runs, x, y, w, h, size=18, color=INK, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    for index, item in enumerate(runs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item[0]
        paragraph.space_after = Pt(6)
        run = paragraph.runs[0]
        run.font.name = font
        run.font.size = Pt(item[1] if len(item) > 1 else size)
        run.font.bold = item[2] if len(item) > 2 else False
        run.font.color.rgb = item[3] if len(item) > 3 else color
    return box


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def footer(slide, number, source="LinkMaker evaluation materials"):
    add_text(slide, source, 0.55, 7.12, 8.8, 0.18, 8, MUTED)
    add_text(slide, str(number), 12.35, 7.08, 0.4, 0.22, 9, MUTED, bold=True, align=PP_ALIGN.RIGHT)


def title(slide, heading, kicker=None, dark=False):
    color = WHITE if dark else INK
    if kicker:
        add_text(slide, kicker.upper(), 0.62, 0.38, 5.5, 0.22, 9, TEAL if not dark else RGBColor(133, 220, 211), bold=True)
    add_text(slide, heading, 0.58, 0.72, 12.1, 0.62, 28, color, bold=True, font="Aptos Display")


def metric_card(slide, x, y, w, h, value, label, accent=TEAL, note=None):
    rect(slide, x, y, w, h, WHITE, radius=True, line=LINE)
    rect(slide, x, y, 0.09, h, accent)
    add_text(slide, value, x + 0.24, y + 0.18, w - 0.4, 0.5, 27, INK, bold=True, font="Aptos Display")
    add_text(slide, label, x + 0.25, y + 0.78, w - 0.4, 0.35, 12, MUTED, bold=True)
    if note:
        add_text(slide, note, x + 0.25, y + 1.14, w - 0.4, 0.35, 9, MUTED)


# 1. Cover
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, NAVY)
rect(slide, 0, 6.82, 13.333, 0.68, TEAL)
add_text(slide, "LINKMAKER", 0.72, 0.74, 4.5, 0.3, 12, RGBColor(133, 220, 211), bold=True)
add_text(slide, "Evidence-aware\nnotebook linking", 0.68, 1.45, 8.6, 1.65, 39, WHITE, bold=True, font="Aptos Display")
add_text(slide, "GPT-5.6 Luna pilot + Markdown/code evaluation", 0.74, 3.45, 8.4, 0.42, 20, RGBColor(212, 225, 234))
add_text(slide, "Titanic notebook case study", 0.74, 4.02, 5.2, 0.3, 13, RGBColor(133, 220, 211), bold=True)
rect(slide, 9.45, 1.55, 2.65, 3.4, RGBColor(27, 48, 71), radius=True, line=RGBColor(63, 91, 116))
add_text(slide, "2", 9.83, 2.0, 1.9, 0.9, 58, WHITE, bold=True, font="Aptos Display", align=PP_ALIGN.CENTER)
add_text(slide, "relationships\nreported", 9.82, 3.02, 1.9, 0.7, 15, RGBColor(212, 225, 234), bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "24 September 2026", 0.74, 6.98, 3.5, 0.2, 10, NAVY, bold=True)

# 2. Run snapshot
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, PALE)
title(slide, "A small, affordable pilot run", "Run snapshot")
add_text(slide, "The Luna run was intentionally compact: enough to demonstrate the workflow and produce two usable links for discussion.", 0.64, 1.48, 11.7, 0.42, 16, MUTED)
metric_card(slide, 0.65, 2.25, 2.75, 1.65, "GPT-5.6", "model", TEAL, "Luna")
metric_card(slide, 3.65, 2.25, 2.75, 1.65, "2", "relationships", ORANGE, "user-reported run result")
metric_card(slide, 6.65, 2.25, 2.75, 1.65, "9k", "input tokens", TEAL, "approximately")
metric_card(slide, 9.65, 2.25, 2.75, 1.65, "1k", "output tokens", ORANGE, "approximately")
rect(slide, 0.65, 4.45, 5.85, 1.45, NAVY, radius=True)
add_text(slide, "Estimated API cost", 0.95, 4.72, 2.8, 0.3, 14, RGBColor(172, 208, 218), bold=True)
add_text(slide, "~$0.030", 0.95, 5.02, 2.4, 0.55, 30, WHITE, bold=True, font="Aptos Display")
add_text(slide, "9,000 x $2/M + 1,000 x $12/M", 3.25, 5.08, 2.8, 0.3, 12, RGBColor(212, 225, 234))
rect(slide, 6.8, 4.45, 5.6, 1.45, WHITE, radius=True, line=LINE)
add_text(slide, "Interpretation", 7.1, 4.72, 2.0, 0.3, 14, TEAL, bold=True)
add_text(slide, "A low-cost feasibility signal, not a statistical performance claim.", 7.1, 5.08, 4.75, 0.45, 17, INK, bold=True)
add_text(slide, "Estimate assumes $2/M input and $12/M output pricing; verify against the active account pricing.", 0.68, 6.33, 11.7, 0.3, 10, MUTED)
footer(slide, 2, "User-provided token counts + repository configuration")

# 3. Pipeline
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, WHITE)
title(slide, "The evidence pipeline", "What changed")
add_text(slide, "Output text gives the linker a second evidence channel while preserving source-code targets.", 0.64, 1.48, 11.7, 0.42, 16, MUTED)
for x, heading, body, accent in [
    (0.8, "1  Markdown claim", "The notebook says what should be true or what the analysis is doing.", TEAL),
    (4.55, "2  Code + output", "Source explains the operation; printed tables and rendered results expose observed values.", ORANGE),
    (8.3, "3  Validated link", "The result stores exact Markdown/code spans and, for outputs, localized evidence.", TEAL),
]:
    rect(slide, x, 2.4, 3.25, 2.3, PALE_BLUE if accent == TEAL else RGBColor(252, 243, 235), radius=True, line=LINE)
    rect(slide, x, 2.4, 3.25, 0.12, accent)
    add_text(slide, heading, x + 0.22, 2.73, 2.8, 0.42, 16, INK, bold=True)
    add_text(slide, body, x + 0.22, 3.35, 2.78, 0.95, 15, MUTED)
    add_text(slide, "evidence", x + 0.22, 4.28, 2.5, 0.22, 10, accent, bold=True)
add_text(slide, "Important guardrail", 0.8, 5.45, 2.4, 0.3, 14, NAVY, bold=True)
add_text(slide, "Execution output supports interpretation; it does not become the code target. Saved output may also be stale, so the prompt favors captured evidence when conflicts appear.", 0.8, 5.82, 11.4, 0.55, 17, INK)
footer(slide, 3, "README.md + routes.py output-context implementation")

# 4. Luna result
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, PALE)
title(slide, "What the Luna result linked", "Representative saved relationship")
rect(slide, 0.68, 1.55, 5.85, 4.85, WHITE, radius=True, line=LINE)
add_text(slide, "Markdown cell 25", 0.98, 1.88, 2.4, 0.3, 12, TEAL, bold=True)
add_text(slide, "- Pclass=3 had most passengers, however most did not survive.\n  Confirms our classifying assumption #2.", 0.98, 2.35, 5.1, 1.1, 20, INK, bold=True)
add_text(slide, "linked to", 0.98, 3.72, 1.2, 0.22, 11, MUTED, bold=True)
add_text(slide, "output from code cell 26", 0.98, 4.08, 3.6, 0.32, 17, NAVY, bold=True)
add_text(slide, "Two bottom-row Pclass=3 histogram groups, split by survival status.", 0.98, 4.7, 4.8, 0.6, 15, MUTED)
add_text(slide, "confidence 0.84", 0.98, 5.65, 2.0, 0.25, 12, ORANGE, bold=True)
rect(slide, 7.0, 1.55, 5.65, 4.85, NAVY, radius=True)
add_text(slide, "Localized evidence", 7.35, 1.9, 3.4, 0.35, 16, RGBColor(172, 208, 218), bold=True)
rect(slide, 7.55, 2.62, 2.0, 2.15, RGBColor(61, 87, 109), radius=True, line=RGBColor(101, 130, 148))
rect(slide, 10.0, 2.62, 2.0, 2.15, RGBColor(61, 87, 109), radius=True, line=RGBColor(101, 130, 148))
for x in (7.78, 10.23):
    rect(slide, x, 3.75, 1.55, 0.55, ORANGE)
add_text(slide, "Pclass=3", 7.88, 2.83, 1.4, 0.24, 11, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "survival split", 10.27, 2.83, 1.45, 0.24, 11, WHITE, bold=True, align=PP_ALIGN.CENTER)
add_text(slide, "2 tight rectangles\nnormalized 0-1000", 7.38, 5.25, 2.8, 0.6, 15, WHITE, bold=True)
add_text(slide, "The user-reported run produced two relationships; this saved artifact contains the representative example shown here.", 0.72, 6.65, 11.8, 0.3, 10, MUTED)
footer(slide, 4, "titanic.linkmaker.json")

# 5. Context evaluation
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, WHITE)
title(slide, "What output text adds", "Recent Markdown/code evaluation")
metric_card(slide, 0.7, 1.55, 2.7, 1.55, "49 / 52", "code cells with text", TEAL, "94.2% of code cells")
metric_card(slide, 3.65, 1.55, 2.7, 1.55, "14,641", "extracted characters", ORANGE, "Titanic notebook")
metric_card(slide, 6.6, 1.55, 2.7, 1.55, "+34.2%", "prompt growth", TEAL, "with output fields")
metric_card(slide, 9.55, 1.55, 2.7, 1.55, "17", "tests passed", ORANGE, "software compatibility")
rect(slide, 0.7, 3.65, 5.85, 2.25, PALE, radius=True, line=LINE)
add_text(slide, "Concrete benefit", 1.0, 3.98, 2.4, 0.3, 14, TEAL, bold=True)
add_text(slide, "The Titanic model-summary output exposes actual scores, including 86.76 for both Random Forest and Decision Tree.", 1.0, 4.38, 4.95, 0.75, 20, INK, bold=True)
add_text(slide, "That evidence is unavailable from variable names alone.", 1.0, 5.35, 4.9, 0.25, 12, MUTED)
rect(slide, 6.8, 3.65, 5.6, 2.25, RGBColor(252, 243, 235), radius=True, line=LINE)
add_text(slide, "Current limitation", 7.1, 3.98, 2.5, 0.3, 14, ORANGE, bold=True)
add_text(slide, "The review did not run a live accuracy comparison. More context can increase tokens, latency, warnings, or stale-output risk.", 7.1, 4.38, 4.75, 0.95, 18, INK, bold=True)
add_text(slide, "Evidence quality improved; model accuracy remains unproven.", 0.72, 6.38, 11.6, 0.3, 14, NAVY, bold=True)
footer(slide, 5, "evaluation/output_context_review.md")

# 6. Benchmark table
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, PALE)
title(slide, "Archived benchmark: tradeoffs are real", "Markdown-to-code evaluation")
add_text(slide, "Frozen Titanic benchmark: 36 unique ground-truth Markdown/code pairs. Metrics are pooled across archived runs.", 0.66, 1.42, 11.8, 0.35, 14, MUTED)
headers = [(0.8, "Configuration"), (4.95, "Precision"), (6.65, "Recall"), (8.25, "F1"), (9.55, "Readout")]
for x, text in headers:
    add_text(slide, text, x, 2.08, 1.7 if x < 9 else 3.0, 0.25, 11, TEAL, bold=True)
rows = [
    ("GPT-5 mini, low", "71.2%", "73.1%", "72.1%", "best balance"),
    ("GPT-5.4 mini, low", "53.0%", "88.9%", "66.4%", "highest recall"),
    ("GPT-5 nano, low", "83.3%", "13.9%", "23.8%", "precise, sparse"),
    ("Current saved file", "64.9%", "66.7%", "65.8%", "single run"),
]
for index, row in enumerate(rows):
    y = 2.45 + index * 0.72
    rect(slide, 0.72, y - 0.1, 11.9, 0.56, WHITE if index % 2 == 0 else RGBColor(245, 249, 248), line=LINE)
    add_text(slide, row[0], 0.9, y + 0.04, 3.7, 0.25, 14, INK, bold=index == 0)
    add_text(slide, row[1], 5.05, y + 0.04, 1.1, 0.25, 14, INK, bold=True)
    add_text(slide, row[2], 6.75, y + 0.04, 1.1, 0.25, 14, INK, bold=True)
    add_text(slide, row[3], 8.35, y + 0.04, 1.1, 0.25, 14, ORANGE if index == 0 else INK, bold=True)
    add_text(slide, row[4], 9.65, y + 0.04, 2.2, 0.25, 13, MUTED)
rect(slide, 0.72, 5.55, 11.9, 0.85, NAVY, radius=True)
add_text(slide, "Takeaway", 1.0, 5.8, 1.2, 0.25, 12, RGBColor(172, 208, 218), bold=True)
add_text(slide, "More links can raise recall while lowering precision. The output-context change still needs a controlled repeated comparison.", 2.25, 5.75, 9.9, 0.35, 16, WHITE, bold=True)
footer(slide, 6, "evaluation/output_context_review.md + archived run JSON")

# 7. Cost and linking quality comparison
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, WHITE)
title(slide, "Cost and linking quality", "Additional model comparison")
add_text(slide, "The two reported runs point to a useful presentation question: how much evidence does each dollar buy?", 0.66, 1.42, 11.8, 0.35, 14, MUTED)
headers = [(0.82, "Model"), (4.25, "Input / output"), (6.25, "Estimated cost"), (8.25, "Observed linking result"), (10.95, "Task")]
for x, text in headers:
    add_text(slide, text, x, 2.08, 2.2, 0.25, 11, TEAL, bold=True)
comparison_rows = [
    ("GPT-5.4-mini-2026-03-17", "9,000 / 500", "$0.009", "1 relationship", "reported run"),
    ("GPT-5.6-terra", "9,000 / 2,000", "$0.042", "count not supplied", "output-chart linking"),
    ("GPT-5.6-luna", "9,000 / 1,000", "~$0.030", "2 relationships", "reported pilot"),
]
for index, row in enumerate(comparison_rows):
    y = 2.46 + index * 0.82
    rect(slide, 0.72, y - 0.1, 11.9, 0.66, PALE if index == 2 else (PALE_BLUE if index == 0 else RGBColor(252, 243, 235)), line=LINE)
    add_text(slide, row[0], 0.9, y + 0.05, 3.05, 0.35, 13, INK, bold=True)
    add_text(slide, row[1], 4.35, y + 0.05, 1.45, 0.3, 13, INK)
    add_text(slide, row[2], 6.35, y + 0.05, 1.45, 0.3, 14, ORANGE, bold=True)
    add_text(slide, row[3], 8.35, y + 0.05, 2.1, 0.3, 13, INK, bold=True)
    add_text(slide, row[4], 11.0, y + 0.05, 1.3, 0.3, 11, MUTED)
rect(slide, 0.72, 5.35, 11.9, 1.0, NAVY, radius=True)
add_text(slide, "Quality reading", 1.0, 5.62, 1.5, 0.25, 12, RGBColor(172, 208, 218), bold=True)
add_text(slide, "GPT-5.4-mini is the cheapest but returned only one relationship. Terra has a higher measured cost for the output-chart task; its relationship count must be added before making a direct yield comparison.", 2.7, 5.55, 9.35, 0.48, 15, WHITE, bold=True)
add_text(slide, "Costs use the token prices supplied for GPT-5.4-mini and the earlier estimate retained for Terra and Luna.", 0.72, 6.62, 11.7, 0.25, 10, MUTED)
footer(slide, 7, "User-provided model run summaries")

# 8. Conclusion
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, NAVY)
title(slide, "The presentation takeaway", "Conclusion", dark=True)
add_text(slide, "A practical proof of concept", 0.72, 1.7, 5.8, 0.35, 18, RGBColor(133, 220, 211), bold=True)
add_text(slide, "Luna produced a small set of output-grounded links at low estimated cost. The links are inspectable, localized, and useful for showing how notebook prose relates to rendered evidence.", 0.72, 2.18, 5.65, 1.35, 22, WHITE, bold=True)
add_text(slide, "A research question remains", 6.95, 1.7, 5.4, 0.35, 18, RGBColor(255, 184, 123), bold=True)
add_text(slide, "Execution text clearly enriches the evidence available to the model. It has not yet been shown to improve precision, recall, or F1 on a controlled benchmark.", 6.95, 2.18, 5.55, 1.35, 22, WHITE, bold=True)
rect(slide, 0.72, 4.45, 11.75, 1.15, RGBColor(27, 48, 71), radius=True, line=RGBColor(63, 91, 116))
add_text(slide, "Next experiment", 1.02, 4.72, 1.8, 0.25, 13, RGBColor(172, 208, 218), bold=True)
add_text(slide, "Run source-only vs output-aware conditions with the same notebook, model, reasoning, validation, and at least three repeats.", 2.85, 4.68, 8.95, 0.35, 17, WHITE, bold=True)
add_text(slide, "Measure precision, recall, F1, invalid spans, latency, input/output tokens, and run-to-run variation.", 1.02, 5.93, 10.9, 0.3, 13, RGBColor(212, 225, 234))
footer(slide, 8, "Recommended next step from the evaluation review")

# 9. Sources
slide = prs.slides.add_slide(blank)
rect(slide, 0, 0, 13.333, 7.5, WHITE)
title(slide, "Sources and assumptions", "Appendix")
source_lines = [
    ("Run snapshot", "User-provided summary: GPT-5.6 Luna, approximately 9k input tokens, 1k output tokens, two relationships."),
    ("Saved artifact", "titanic.linkmaker.json records outputModel = gpt-5.6-luna and a representative output-grounded relationship."),
    ("Implementation", "README.md and LinkMaker/routes.py describe output text extraction, source targeting, and output localization."),
    ("Evaluation", "evaluation/output_context_review.md, evaluation/titanic_link_review.md, and archived run JSON files."),
    ("Cost", "Illustrative estimate using $2 per million input tokens and $12 per million output tokens: $0.018 + $0.012 = $0.030."),
]
for index, (label, text) in enumerate(source_lines):
    y = 1.55 + index * 0.86
    rect(slide, 0.7, y, 1.8, 0.42, PALE, radius=True, line=LINE)
    add_text(slide, label, 0.86, y + 0.1, 1.45, 0.2, 11, TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, text, 2.85, y + 0.03, 9.25, 0.48, 15, INK)
rect(slide, 0.7, 6.15, 11.75, 0.55, RGBColor(252, 243, 235), radius=True, line=LINE)
add_text(slide, "The deck presents the Luna run as a feasibility result and does not claim a causal accuracy improvement.", 1.0, 6.32, 11.1, 0.22, 13, ORANGE, bold=True, align=PP_ALIGN.CENTER)
footer(slide, 9, "Prepared from the current workspace")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(OUT)
